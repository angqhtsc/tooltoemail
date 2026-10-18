"""ToolToEmail v2 - send each student's result file to the parents' email(s), from an Excel list."""
import configparser
import datetime
import html as htmllib
import json
import mimetypes
import os
import queue
import re
import smtplib
import socket
import sqlite3
import ssl
import subprocess
import sys
import threading
import tkinter as tk
from email.message import EmailMessage
from email.utils import formataddr
from tkinter import filedialog, messagebox, ttk

from openpyxl import Workbook, load_workbook

IS_MAC = sys.platform == "darwin"
IS_WIN = sys.platform.startswith("win")
if IS_MAC:  # data must live outside the .app bundle (it may be read-only / replaced on update)
    APP_DIR = os.path.join(os.path.expanduser("~"), "ToolToEmail")
    os.makedirs(APP_DIR, exist_ok=True)
else:
    APP_DIR = os.path.dirname(sys.executable if getattr(sys, "frozen", False) else os.path.abspath(__file__))
CONFIG_PATH = os.path.join(APP_DIR, "config.ini")
SETTINGS_PATH = os.path.join(APP_DIR, "last_paths.json")
DB_PATH = os.path.join(APP_DIR, "send_log.db")

EMAIL_RE = re.compile(r"^[^@\s,;]+@[^@\s,;]+\.[^@\s,;]+$")
MAX_ATTACH_MB = 24
DONE = ("sent", "partial")  # statuses that count as "already delivered"
TICKER_TEXT = "From K.Chi C1.3 with <3"  # running text at the bottom of the window
ALL_SUBJECTS = "(all subjects)"

DEFAULT_CONFIG = """; ToolToEmail settings. Lines starting with ; are comments.
; The email TEXT is NOT written here - it lives in the file named by body_file (default body.txt).

[account]
email =
password =
sender_name = School Office
; provider: auto | gmail | yahoo | icloud | zoho | custom   (auto = guess from your email address)
provider = auto
; Only needed for provider = custom (school / company / hosting mail). Ask your mail provider or IT:
smtp_host =
smtp_port = 587
; security: ssl (usually port 465) or starttls (usually port 587)
security = starttls
; username: leave empty if it is the same as your email address
username =

[email]
; You can use {name}, {id} and any other column title of your Excel file, e.g. {class}
subject = Exam result of {name}
; body.txt = plain text.  Use a name ending in .html (e.g. body_example.html) for a formatted email.
body_file = body.txt

[sending]
; seconds to wait between two emails
delay_seconds = 2
"""

DEFAULT_BODY = """Dear Parent of {name},

Please find attached the result of {name} (ID {id}).

Best regards,
School Office
"""

PRESETS = {  # provider -> (host, port, security)
    "gmail": ("smtp.gmail.com", 465, "ssl"),
    "yahoo": ("smtp.mail.yahoo.com", 465, "ssl"),
    "icloud": ("smtp.mail.me.com", 587, "starttls"),
    "zoho": ("smtp.zoho.com", 465, "ssl"),
}
DOMAINS = {"gmail.com": "gmail", "googlemail.com": "gmail", "ymail.com": "yahoo", "rocketmail.com": "yahoo",
           "icloud.com": "icloud", "me.com": "icloud", "mac.com": "icloud", "zoho.com": "zoho", "zohomail.com": "zoho"}
MS_DOMAINS = {"outlook.com", "hotmail.com", "live.com", "msn.com"}
AUTH_HINT = {
    "gmail": "Gmail needs an App password (16 letters from myaccount.google.com/apppasswords), not your normal "
             "password and not backup codes.",
    "yahoo": "Yahoo needs an App password (Yahoo Account Security > Generate app password).",
    "icloud": "iCloud needs an app-specific password (appleid.apple.com > Sign-In and Security).",
    "zoho": "Zoho needs an application-specific password if two-factor login is on.",
}

HEADERS = {
    "id": {"id", "studentid", "number", "no", "studentnumber", "code", "stt"},
    "name": {"name", "studentname", "fullname", "student"},
    "e1": {"parentemail1", "email1", "parentemail", "email", "emailaddress1", "parent1email", "emailparent1"},
    "e2": {"parentemail2", "email2", "emailaddress2", "parent2email", "emailparent2"},
}


class ConfigError(Exception):
    pass


# --------------------------------------------------------------------------- config / settings / log db
def read_cp():
    cp = configparser.ConfigParser(interpolation=None, strict=False)
    try:
        cp.read(CONFIG_PATH, encoding="utf-8-sig")
    except (configparser.Error, UnicodeDecodeError) as e:
        raise ConfigError("config.ini cannot be read:\n\n" + str(e) + "\n\nTip: keep config.ini for short settings only. "
                          "The email text belongs in body.txt. Delete config.ini to get a fresh copy.")
    return cp


def ensure_files():
    """Create config.ini / body.txt when missing (an old long body inside config.ini is moved to body.txt)."""
    if not os.path.exists(CONFIG_PATH):
        with open(CONFIG_PATH, "w", encoding="utf-8") as f:
            f.write(DEFAULT_CONFIG)
    body = os.path.join(APP_DIR, "body.txt")
    if not os.path.exists(body):
        text = DEFAULT_BODY
        try:
            cp = read_cp()
            if cp.has_option("email", "body") and cp.get("email", "body").strip():
                text = cp.get("email", "body").strip() + "\n"
        except ConfigError:
            pass
        with open(body, "w", encoding="utf-8") as f:
            f.write(text)


def read_text(path):
    for enc in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            with open(path, encoding=enc) as f:
                return f.read()
        except (UnicodeError, UnicodeDecodeError):
            continue
    with open(path, encoding="utf-8", errors="replace") as f:
        return f.read()


def load_config():
    ensure_files()
    cp = read_cp()

    def g(key, sections, default=""):
        for sec in sections:
            if cp.has_option(sec, key) and cp.get(sec, key).strip():
                return cp.get(sec, key).strip()
        return default

    acct = ("account", "gmail")  # [gmail] = settings file of the first version
    try:
        delay = float(g("delay_seconds", ("sending",), "2"))
    except ValueError:
        delay = 2.0
    body_file = g("body_file", ("email",), "body.txt")
    body_path = body_file if os.path.isabs(body_file) else os.path.join(APP_DIR, body_file)
    if not os.path.isfile(body_path):
        raise ConfigError(f"The email body file was not found:\n{body_path}\n\nCheck 'body_file' in config.ini.")
    return {
        "email": g("email", acct), "password": g("password", acct).replace(" ", ""),
        "sender_name": g("sender_name", acct), "provider": g("provider", acct, "auto").lower(),
        "smtp_host": g("smtp_host", acct), "smtp_port": g("smtp_port", acct), "security": g("security", acct).lower(),
        "username": g("username", acct),
        "subject": g("subject", ("email",), "Result of {name}"),
        "body": read_text(body_path), "body_path": body_path,
        "is_html": body_path.lower().endswith((".html", ".htm")), "delay": delay,
    }


def load_settings():
    try:
        with open(SETTINGS_PATH, encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}


def save_settings(data):
    try:
        with open(SETTINGS_PATH, "w", encoding="utf-8") as f:
            json.dump(data, f)
    except OSError:
        pass


def db(sql, args=()):
    con = sqlite3.connect(DB_PATH, timeout=10)
    try:
        rows = con.execute(sql, args).fetchall()
        con.commit()
        return rows
    finally:
        con.close()


def init_db():
    db("""CREATE TABLE IF NOT EXISTS log(id INTEGER PRIMARY KEY AUTOINCREMENT, ts TEXT, student_id TEXT, name TEXT,
          recipients TEXT, attachments TEXT, status TEXT, error TEXT, subject TEXT DEFAULT '', batch TEXT DEFAULT '')""")
    have = [r[1] for r in db("PRAGMA table_info(log)")]
    for col in ("subject", "batch"):  # logs made by the first version
        if col not in have:
            db(f"ALTER TABLE log ADD COLUMN {col} TEXT DEFAULT ''")


def log_row(student, status, error, subject, batch):
    db("INSERT INTO log(ts,student_id,name,recipients,attachments,status,error,subject,batch) VALUES(?,?,?,?,?,?,?,?,?)",
       (datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"), student["id"], student["name"],
        ", ".join(student["emails"]), ", ".join(os.path.basename(p) for p in student["files"]), status, error,
        subject, batch))


def last_status(student_id, batch):
    r = db("SELECT status FROM log WHERE student_id=? AND batch=? AND status!='test' ORDER BY id DESC LIMIT 1",
           (student_id, batch))
    return r[0][0] if r else ""


# --------------------------------------------------------------------------- reading Excel + matching files
def norm(s):
    return re.sub(r"[^a-z0-9]", "", str(s).lower())


def cell_text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def id_key(s):
    s = s.strip().lower()
    return (s.lstrip("0") or "0") if s.isdigit() else s


def load_students(xlsx_path, folder, batch):
    """Return (students, unmatched_files). 'batch' = subject template; sent-status is tracked per batch."""
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise ValueError("The Excel file is empty.")
    header = [cell_text(h) for h in rows[0]]
    cols = {}
    for idx, h in enumerate(header):
        for field, alts in HEADERS.items():
            if norm(h) in alts and field not in cols:
                cols[field] = idx
    if "id" not in cols or "e1" not in cols:
        raise ValueError("Could not find the required columns in row 1.\nNeeded: ID, Name, Parent Email 1 "
                         "(and optionally Parent Email 2).\nFound: " + ", ".join(h for h in header if h))
    files = {}
    for fn in sorted(os.listdir(folder)):
        full = os.path.join(folder, fn)
        if not os.path.isfile(full) or fn.startswith(("~$", ".")):
            continue
        token = re.split(r"[_\-\s.]", os.path.splitext(fn)[0])[0]
        files.setdefault(id_key(token), []).append(full)

    students, seen, used = [], {}, set()
    for rownum, r in enumerate(rows[1:], 2):
        get = lambda f: cell_text(r[cols[f]]) if f in cols and cols[f] < len(r) else ""
        sid, name = get("id"), get("name")
        emails = list(dict.fromkeys(e for e in re.split(r"[;,\s]+", get("e1") + " " + get("e2")) if e))
        if not (sid or name or emails):
            continue
        raw = {header[i]: cell_text(v) for i, v in enumerate(r) if i < len(header) and header[i]}
        problems = []
        if not sid:
            problems.append("Missing ID")
        elif id_key(sid) in seen:
            problems.append(f"Duplicate ID (also row {seen[id_key(sid)]})")
        seen.setdefault(id_key(sid), rownum)
        if not emails:
            problems.append("No email")
        bad = [e for e in emails if not EMAIL_RE.match(e)]
        if bad:
            problems.append("Invalid email: " + ", ".join(bad))
        matched = files.get(id_key(sid), []) if sid else []
        used.add(id_key(sid))
        if not matched:
            problems.append("No result file found")
        elif sum(os.path.getsize(p) for p in matched) > MAX_ATTACH_MB * 1024 * 1024:
            problems.append(f"Attachments larger than {MAX_ATTACH_MB} MB")
        students.append({"id": sid, "name": name, "emails": emails, "files": matched, "raw": raw, "problems": problems,
                         "status": last_status(sid, batch) if sid else ""})
    dup = {id_key(s["id"]) for s in students if any(p.startswith("Duplicate ID") for p in s["problems"])}
    for s in students:  # block the first occurrence too: we cannot know which row is right
        if s["id"] and id_key(s["id"]) in dup and not any(p.startswith("Duplicate ID") for p in s["problems"]):
            s["problems"].append("Duplicate ID (also in a later row)")
    unmatched = [os.path.basename(p) for k, ps in files.items() if k not in used for p in ps]
    return students, unmatched


# --------------------------------------------------------------------------- building / sending emails
def fill(template, values, escape=False):
    for k, v in values.items():
        template = template.replace("{" + k + "}", htmllib.escape(v) if escape else v)
    return template


def html_to_text(h):
    h = re.sub(r"(?is)<(script|style).*?>.*?</\1>", "", h)
    h = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</tr>|</h\d>|</li>", "\n", h)
    h = htmllib.unescape(re.sub(r"<[^>]+>", "", h))
    return re.sub(r"\n{3,}", "\n\n", h).strip()


def values_of(s):
    v = dict(s["raw"])
    v.update(id=s["id"], name=s["name"])
    return v


def render_subject(cfg, s, test=False):
    return ("[TEST] " if test else "") + " ".join(fill(cfg["subject"], values_of(s)).split())


def build_message(cfg, s, test_to=None):
    values = values_of(s)
    msg = EmailMessage()
    msg["Subject"] = render_subject(cfg, s, bool(test_to))
    msg["From"] = formataddr((cfg["sender_name"], cfg["email"]))
    msg["To"] = test_to or ", ".join(s["emails"])
    if cfg["is_html"]:
        page = fill(cfg["body"], values, escape=True)
        msg.set_content(html_to_text(page))
        msg.add_alternative(page, subtype="html")
    else:
        msg.set_content(fill(cfg["body"], values).rstrip() + "\n")
    for path in s["files"]:
        ctype = mimetypes.guess_type(path)[0] or "application/octet-stream"
        maintype, subtype = ctype.split("/", 1)
        with open(path, "rb") as f:
            msg.add_attachment(f.read(), maintype=maintype, subtype=subtype, filename=os.path.basename(path))
    return msg


def resolve_server(cfg):
    """Return (host, port, security, provider) from config; raises ConfigError with a clear message."""
    domain = cfg["email"].split("@")[-1].lower()
    provider = cfg["provider"] or "auto"
    if provider == "auto":
        provider = DOMAINS.get(domain) or ("yahoo" if domain.startswith("yahoo.") else
                                           "microsoft" if domain in MS_DOMAINS else "custom")
    if cfg["smtp_host"]:
        try:
            port = int(cfg["smtp_port"] or 587)
        except ValueError:
            raise ConfigError("smtp_port in config.ini must be a number (465 or 587).")
        sec = {"tls": "starttls", "": "ssl" if port == 465 else "starttls"}.get(cfg["security"], cfg["security"])
        if sec not in ("ssl", "starttls", "none"):
            raise ConfigError("security in config.ini must be ssl, starttls or none.")
        return cfg["smtp_host"], port, sec, provider
    if provider in PRESETS:
        return PRESETS[provider] + (provider,)
    if provider == "microsoft":
        raise ConfigError("Outlook / Hotmail / Live accounts cannot be used any more: Microsoft switched off "
                          "password (and app password) sign-in for sending email. Please use a Gmail, Yahoo, iCloud, "
                          "Zoho or school/company mail account instead.")
    raise ConfigError("Your email provider is not one of the built-in ones.\nPlease fill in smtp_host, smtp_port and "
                      "security in config.ini (ask your mail provider or IT for the 'SMTP server' settings).")


def connect(cfg):
    host, port, sec, _ = resolve_server(cfg)
    ctx = ssl.create_default_context()
    if sec == "ssl":
        server = smtplib.SMTP_SSL(host, port, context=ctx, timeout=60)
    else:
        server = smtplib.SMTP(host, port, timeout=60)
        server.ehlo()
        if sec == "starttls":
            server.starttls(context=ctx)
            server.ehlo()
    server.login(cfg["username"] or cfg["email"], cfg["password"])
    return server


def explain_error(cfg, e):
    if isinstance(e, ConfigError):
        return str(e)
    if isinstance(e, smtplib.SMTPAuthenticationError):
        try:
            hint = AUTH_HINT.get(resolve_server(cfg)[3], "Check the username and password, and make sure your mail "
                                 "provider allows SMTP sending from apps.")
        except ConfigError:
            hint = ""
        return f"The mail server refused the login.\n\n{hint}\n\nServer said: {e}"
    return f"Cannot connect to the mail server: {type(e).__name__}: {e}"


# --------------------------------------------------------------------------- GUI
STATUS_TEXT = {"": "not sent", "sent": "SENT", "failed": "FAILED", "partial": "SENT (check log)"}


def open_path(path):
    """Open a file or folder with the default program (Windows / macOS / Linux)."""
    try:
        if IS_WIN:
            os.startfile(path)
        else:
            subprocess.Popen(["open" if IS_MAC else "xdg-open", path])
    except Exception:
        messagebox.showinfo("File", path)


def open_text(path):
    """Open a text file in a plain-text editor (Notepad / TextEdit)."""
    try:
        if IS_WIN:
            subprocess.Popen(["notepad.exe", path])
        elif IS_MAC:
            subprocess.Popen(["open", "-t", path])
        else:
            open_path(path)
    except OSError:
        open_path(path)


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("ToolToEmail v2 - send results to parents")
        self.geometry("1200x700")
        init_db()
        try:
            ensure_files()
        except OSError:
            pass
        self.students, self.selected, self.batch = [], set(), ""
        self.q, self.stop, self.busy = queue.Queue(), threading.Event(), False
        st = load_settings()
        self.excel = tk.StringVar(value=st.get("excel", ""))
        self.folder = tk.StringVar(value=st.get("folder", ""))
        self.info = tk.StringVar(value="Choose the Excel file and the results folder, then click 'Load & check'.  "
                                       f"(Settings and log are stored in: {APP_DIR})")
        self.build_ui()
        self.after(200, self.poll)

    def build_ui(self):
        self.ticker = tk.Canvas(self, height=26, highlightthickness=0, bg="#eef3fb")
        self.ticker.pack(side="bottom", fill="x")
        self.ticker_item = self.ticker.create_text(1200, 13, text=TICKER_TEXT, anchor="w", fill="#1F3864",
                                                   font=("Helvetica", 12 if IS_MAC else 10, "italic"))
        self.after(40, self.run_ticker)
        nb = ttk.Notebook(self)
        nb.pack(fill="both", expand=True)
        send_tab, log_tab = ttk.Frame(nb), ttk.Frame(nb)
        nb.add(send_tab, text="  Send  ")
        nb.add(log_tab, text="  Log  ")
        nb.bind("<<NotebookTabChanged>>", lambda e: self.refresh_log())

        top = ttk.Frame(send_tab, padding=8)
        top.pack(fill="x")
        for r, (label, var, cmd) in enumerate([("Excel file:", self.excel, self.pick_excel),
                                               ("Results folder:", self.folder, self.pick_folder)]):
            ttk.Label(top, text=label, width=14).grid(row=r, column=0, sticky="w", pady=2)
            ttk.Entry(top, textvariable=var).grid(row=r, column=1, sticky="we", padx=4)
            ttk.Button(top, text="Browse...", command=cmd).grid(row=r, column=2)
        top.columnconfigure(1, weight=1)
        ttk.Button(top, text="Load & check", command=self.load).grid(row=0, column=3, rowspan=2, padx=8, sticky="ns")
        tb = ttk.Frame(send_tab, padding=(8, 0))
        tb.pack(fill="x")
        for text, cmd in [("Edit email body", self.edit_body), ("Open config.ini", self.open_config),
                          ("Open data folder", lambda: open_path(APP_DIR)),
                          ("Preview email", self.preview), ("Test login", self.test_login)]:
            ttk.Button(tb, text=text, command=cmd).pack(side="left", padx=(0, 4))

        cols = [("sel", "", 40), ("id", "ID", 70), ("name", "Student", 190), ("emails", "Parent email(s)", 280),
                ("files", "Attachment(s)", 190), ("check", "Check", 200), ("status", "Status", 110)]
        frame = ttk.Frame(send_tab)
        frame.pack(fill="both", expand=True, padx=8, pady=(6, 0))
        self.tree = ttk.Treeview(frame, columns=[c[0] for c in cols], show="headings", selectmode="browse")
        for key, text, w in cols:
            self.tree.heading(key, text=text)
            self.tree.column(key, width=w, anchor="center" if key == "sel" else "w")
        sb = ttk.Scrollbar(frame, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=sb.set)
        self.tree.pack(side="left", fill="both", expand=True)
        sb.pack(side="right", fill="y")
        self.tree.tag_configure("bad", foreground="#b00020")
        self.tree.tag_configure("done", foreground="#1b7a2f")
        self.tree.bind("<Button-1>", self.on_click)
        self.tree.bind("<Double-Button-1>", self.open_attachment)

        bar = ttk.Frame(send_tab, padding=8)
        bar.pack(fill="x")
        ttk.Button(bar, text="Select not-yet-sent", command=lambda: self.select(lambda s: s["status"] not in DONE)).pack(side="left")
        ttk.Button(bar, text="Select all OK", command=lambda: self.select(lambda s: True)).pack(side="left", padx=4)
        ttk.Button(bar, text="Select failed only", command=lambda: self.select(lambda s: s["status"] == "failed")).pack(side="left")
        ttk.Button(bar, text="Deselect all", command=lambda: self.select(lambda s: False)).pack(side="left", padx=4)
        self.btn_stop = ttk.Button(bar, text="Stop", command=self.stop.set, state="disabled")
        self.btn_stop.pack(side="right")
        self.btn_send = ttk.Button(bar, text="Send selected...", command=self.send_selected)
        self.btn_send.pack(side="right", padx=4)
        self.btn_test = ttk.Button(bar, text="Send test to myself", command=self.send_test)
        self.btn_test.pack(side="right")
        self.progress = ttk.Progressbar(send_tab, mode="determinate")
        self.progress.pack(fill="x", padx=8)
        ttk.Label(send_tab, textvariable=self.info, padding=8, wraplength=1100).pack(fill="x")
        ttk.Label(send_tab, text="Tip: click the checkbox to (de)select a student; double-click a row to open its attachment.",
                  foreground="#666").pack(anchor="w", padx=8, pady=(0, 6))

        ctl = ttk.Frame(log_tab, padding=8)
        ctl.pack(fill="x")
        ttk.Label(ctl, text="Status:").pack(side="left")
        self.log_status = ttk.Combobox(ctl, values=["all", "sent", "failed", "partial", "test"], state="readonly", width=8)
        self.log_status.set("all")
        self.log_status.pack(side="left", padx=(4, 10))
        ttk.Label(ctl, text="Subject:").pack(side="left")
        self.log_subject = ttk.Combobox(ctl, values=[ALL_SUBJECTS], state="readonly", width=40)
        self.log_subject.set(ALL_SUBJECTS)
        self.log_subject.pack(side="left", padx=(4, 10))
        ttk.Label(ctl, text="Search:").pack(side="left")
        self.log_search = tk.StringVar()
        ent = ttk.Entry(ctl, textvariable=self.log_search, width=22)
        ent.pack(side="left", padx=4)
        ent.bind("<KeyRelease>", lambda e: self.refresh_log())
        for cb in (self.log_status, self.log_subject):
            cb.bind("<<ComboboxSelected>>", lambda e: self.refresh_log())
        ttk.Button(ctl, text="Export to Excel...", command=self.export_log).pack(side="right")
        self.log_count = tk.StringVar()
        ttk.Label(log_tab, textvariable=self.log_count, padding=(8, 0)).pack(anchor="w")
        lcols = [("ts", "Time", 130), ("subject", "Subject", 240), ("sid", "ID", 60), ("name", "Student", 150),
                 ("rec", "Recipients", 230), ("st", "Status", 70), ("err", "Error / note", 280)]
        lf = ttk.Frame(log_tab)
        lf.pack(fill="both", expand=True, padx=8, pady=8)
        self.log_tree = ttk.Treeview(lf, columns=[c[0] for c in lcols], show="headings")
        for key, text, w in lcols:
            self.log_tree.heading(key, text=text)
            self.log_tree.column(key, width=w, anchor="w")
        lsb = ttk.Scrollbar(lf, orient="vertical", command=self.log_tree.yview)
        self.log_tree.configure(yscrollcommand=lsb.set)
        self.log_tree.pack(side="left", fill="both", expand=True)
        lsb.pack(side="right", fill="y")
        self.log_tree.tag_configure("failed", foreground="#b00020")

    def run_ticker(self):
        """Scroll the footer text from right to left, forever."""
        self.ticker.move(self.ticker_item, -2, 0)
        x1, _, x2, _ = self.ticker.bbox(self.ticker_item)
        if x2 < 0:  # fully left of the window -> come back in from the right edge
            self.ticker.move(self.ticker_item, max(self.ticker.winfo_width(), 300) - x1, 0)
        self.after(40, self.run_ticker)

    # ---- pickers / files
    def pick_excel(self):
        p = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx *.xlsm"), ("All", "*.*")])
        if p:
            self.excel.set(p)

    def pick_folder(self):
        p = filedialog.askdirectory()
        if p:
            self.folder.set(p)

    def open_config(self):
        try:
            ensure_files()
        except OSError:
            pass
        open_text(CONFIG_PATH)  # never parses the file, so a broken config.ini can still be opened and fixed

    def edit_body(self):
        try:
            ensure_files()
            open_text(load_config()["body_path"])
        except ConfigError as e:
            messagebox.showerror("Email body", str(e))

    def open_attachment(self, _e):
        sel = self.tree.selection()
        if sel and self.students[int(sel[0])]["files"]:
            open_path(self.students[int(sel[0])]["files"][0])

    def get_config(self):
        try:
            return load_config()
        except ConfigError as e:
            if messagebox.askyesno("Problem in the settings", str(e) + "\n\nOpen config.ini now?"):
                open_text(CONFIG_PATH)
            return None

    # ---- table
    def load(self):
        if not os.path.isfile(self.excel.get()) or not os.path.isdir(self.folder.get()):
            messagebox.showwarning("Missing input", "Please choose an existing Excel file and results folder.")
            return
        cfg = self.get_config()
        if not cfg:
            return
        save_settings({"excel": self.excel.get(), "folder": self.folder.get()})
        self.batch = cfg["subject"]
        try:
            self.students, unmatched = load_students(self.excel.get(), self.folder.get(), self.batch)
        except Exception as e:
            messagebox.showerror("Cannot read the Excel file", str(e))
            return
        self.selected = {str(i) for i, s in enumerate(self.students) if not s["problems"] and s["status"] not in DONE}
        self.tree.delete(*self.tree.get_children())
        for i in range(len(self.students)):
            self.tree.insert("", "end", iid=str(i), values=self.row_values(i), tags=self.row_tags(i))
        bad = sum(1 for s in self.students if s["problems"])
        self.update_info(f"Subject: {self.batch} | {len(self.students)} students | {bad} with problems (red, cannot be "
                         f"selected) | {len(unmatched)} result files match no student")
        if unmatched:
            shown = ", ".join(unmatched[:10]) + (" ..." if len(unmatched) > 10 else "")
            messagebox.showinfo("Files without a matching student", f"{len(unmatched)} file(s) in the folder do not "
                                f"match any ID in the Excel file:\n\n{shown}")

    def row_values(self, i):
        s = self.students[i]
        box = "—" if s["problems"] else ("☑" if str(i) in self.selected else "☐")
        return (box, s["id"], s["name"], ", ".join(s["emails"]), ", ".join(os.path.basename(p) for p in s["files"]),
                "; ".join(s["problems"]) or "OK", STATUS_TEXT.get(s["status"], s["status"]))

    def row_tags(self, i):
        s = self.students[i]
        return ("bad",) if s["problems"] else (("done",) if s["status"] in DONE else ())

    def refresh_row(self, i):
        self.tree.item(str(i), values=self.row_values(i), tags=self.row_tags(i))

    def on_click(self, e):
        if self.tree.identify_region(e.x, e.y) == "cell" and self.tree.identify_column(e.x) == "#1":
            iid = self.tree.identify_row(e.y)
            if iid and not self.students[int(iid)]["problems"] and not self.busy:
                self.selected ^= {iid}
                self.refresh_row(int(iid))
                self.update_info()

    def select(self, pred):
        self.selected = {str(i) for i, s in enumerate(self.students) if not s["problems"] and pred(s)}
        for i in range(len(self.students)):
            self.refresh_row(i)
        self.update_info()

    def update_info(self, prefix=None):
        if prefix is not None:
            self._prefix = prefix
        done = sum(1 for s in self.students if s["status"] in DONE)
        self.info.set(f"{getattr(self, '_prefix', '')} | already sent with this subject: {done} | selected now: {len(self.selected)}")

    def preview(self):
        if not self.students:
            messagebox.showinfo("Preview", "Load & check the list first.")
            return
        sel = self.tree.selection()
        idx = int(sel[0]) if sel else next(iter(sorted(map(int, self.selected))), None)
        if idx is None:
            idx = next((i for i, s in enumerate(self.students) if not s["problems"]), 0)
        cfg = self.get_config()
        if not cfg:
            return
        s = self.students[idx]
        try:
            msg = build_message(cfg, s)
            body = msg.get_body(preferencelist=("plain",)).get_content()
        except Exception as e:
            messagebox.showerror("Preview", f"{type(e).__name__}: {e}")
            return
        top = tk.Toplevel(self)
        top.title(f"Preview - ID {s['id']} {s['name']}")
        top.geometry("760x560")
        txt = tk.Text(top, wrap="word", padx=10, pady=10)
        sc = ttk.Scrollbar(top, orient="vertical", command=txt.yview)
        txt.configure(yscrollcommand=sc.set)
        sc.pack(side="right", fill="y")
        txt.pack(fill="both", expand=True)
        att = ", ".join(os.path.basename(p) for p in s["files"]) or "(none)"
        txt.insert("1.0", f"From: {msg['From']}\nTo: {msg['To']}\nSubject: {msg['Subject']}\nAttachments: {att}\n"
                          f"{'-' * 70}\n{body}")
        if cfg["is_html"]:
            txt.insert("end", "\n\n(The real email is formatted HTML; this preview shows its plain-text version.)")
        txt.configure(state="disabled")

    # ---- sending
    def check_credentials(self):
        cfg = self.get_config()
        if not cfg:
            return None
        if not cfg["email"] or not cfg["password"]:
            if messagebox.askyesno("Email account not set up", "Please fill in your email address and password in "
                                   "config.ini.\n\nOpen config.ini now?"):
                open_text(CONFIG_PATH)
            return None
        try:
            resolve_server(cfg)
        except ConfigError as e:
            messagebox.showerror("Email provider", str(e))
            return None
        return cfg

    def test_login(self):
        cfg = self.check_credentials()
        if cfg:
            self.info.set("Testing the login ...")
            threading.Thread(target=self.login_worker, args=(cfg,), daemon=True).start()

    def login_worker(self, cfg):
        try:
            connect(cfg).quit()
            self.q.put(("msg", "Login OK", f"Connected to the mail server and signed in as {cfg['email']}."))
        except Exception as e:
            self.q.put(("fatal", explain_error(cfg, e)))

    def start(self, jobs, cfg, test_to=None):
        self.busy = True
        self.stop.clear()
        for b in (self.btn_send, self.btn_test):
            b.config(state="disabled")
        self.btn_stop.config(state="normal")
        self.progress.config(maximum=len(jobs), value=0)
        threading.Thread(target=self.worker, args=(jobs, cfg, test_to), daemon=True).start()

    def send_selected(self):
        if self.busy:
            return
        jobs = [s for i, s in enumerate(self.students) if str(i) in self.selected]
        if not jobs:
            messagebox.showinfo("Nothing selected", "Select at least one student first.")
            return
        cfg = self.check_credentials()
        if not cfg:
            return
        if cfg["subject"] != self.batch:
            messagebox.showwarning("Subject changed", "The subject in config.ini changed after you loaded the list.\n"
                                   "Click 'Load & check' again so the sent/failed status matches the new subject.")
            return
        again = [s for s in jobs if s["status"] in DONE]
        if again and not messagebox.askyesno("Already sent", f"{len(again)} selected student(s) were already sent with "
                                             "this subject. Send to them AGAIN?\n\nYes = send again, No = skip them."):
            jobs = [s for s in jobs if s["status"] not in DONE]
            if not jobs:
                return
        n_rec = sum(len(s["emails"]) for s in jobs)
        if messagebox.askyesno("Confirm sending", f"Send {len(jobs)} emails ({n_rec} recipients)?\n\nFrom: {cfg['email']}\n"
                               f"Subject example: {render_subject(cfg, jobs[0])}\n\nThis cannot be undone."):
            self.start(jobs, cfg)

    def send_test(self):
        if self.busy:
            return
        pool = [s for i, s in enumerate(self.students) if str(i) in self.selected] or [s for s in self.students if not s["problems"]]
        if not pool:
            messagebox.showinfo("No student", "Load the list first (a student without problems is needed).")
            return
        cfg = self.check_credentials()
        if cfg and messagebox.askyesno("Test email", f"Send a TEST email (with the attachment of {pool[0]['name']}, "
                                       f"ID {pool[0]['id']}) to yourself at {cfg['email']}?"):
            self.start(pool[:1], cfg, test_to=cfg["email"])

    def worker(self, jobs, cfg, test_to=None):
        server, total, batch = None, len(jobs), cfg["subject"]
        try:
            for n, s in enumerate(jobs, 1):
                if self.stop.is_set():
                    break
                status, err, subject = "failed", "", render_subject(cfg, s, bool(test_to))
                for _attempt in (1, 2):
                    try:
                        if server is None:
                            server = connect(cfg)
                        refused = server.send_message(build_message(cfg, s, test_to))
                        status, err = ("test" if test_to else "sent"), ""
                        if refused:
                            status, err = "partial", "Refused: " + ", ".join(refused)
                        break
                    except (smtplib.SMTPAuthenticationError, ConfigError) as e:
                        self.q.put(("fatal", explain_error(cfg, e)))
                        return
                    except smtplib.SMTPRecipientsRefused as e:
                        err = "Recipient refused: " + ", ".join(e.recipients)
                        break
                    except (smtplib.SMTPServerDisconnected, ConnectionError, TimeoutError, ssl.SSLError, socket.gaierror) as e:
                        err, server = f"Connection problem: {e}", None  # retry once on a fresh connection
                    except Exception as e:
                        err = f"{type(e).__name__}: {e}"
                        if any(k in err.lower() for k in ("5.4.5", "daily", "quota", "too many", "sending limit")):
                            log_row(s, "failed", err, subject, batch)
                            self.q.put(("row", s["id"], "failed", err, n, total))
                            self.q.put(("fatal", "The mail provider's sending limit was reached. Try again later "
                                                 "(often tomorrow); only the unsent students remain to be resent.\n\n" + err))
                            return
                        break
                log_row(s, status, err, subject, batch)
                self.q.put(("row", s["id"], status, err, n, total))
                if n < total and self.stop.wait(cfg["delay"]):
                    break
        finally:
            try:
                if server:
                    server.quit()
            except Exception:
                pass
            self.q.put(("done",))

    def poll(self):
        try:
            while True:
                ev = self.q.get_nowait()
                if ev[0] == "row":
                    _, sid, status, err, n, total = ev
                    self.progress.config(value=n)
                    for i, s in enumerate(self.students):
                        if s["id"] == sid and status != "test":
                            s["status"] = status
                            self.selected.discard(str(i))
                            self.refresh_row(i)
                    self.info.set(f"Sending {n}/{total} ... last: ID {sid} -> {status.upper()} {err}")
                    if status == "test":
                        messagebox.showinfo("Test email", "Test email sent. Check your inbox.")
                elif ev[0] == "msg":
                    messagebox.showinfo(ev[1], ev[2])
                elif ev[0] == "fatal":
                    messagebox.showerror("Stopped", ev[1])
                elif ev[0] == "done":
                    self.busy = False
                    self.btn_send.config(state="normal")
                    self.btn_test.config(state="normal")
                    self.btn_stop.config(state="disabled")
                    failed = sum(1 for s in self.students if s["status"] == "failed")
                    self.update_info("Finished" + (f" - {failed} failed: click 'Select failed only' then send again" if failed else ""))
                    self.refresh_log()
        except queue.Empty:
            pass
        self.after(200, self.poll)

    # ---- log tab
    def refresh_log(self):
        batches = [r[0] for r in db("SELECT DISTINCT batch FROM log WHERE batch!='' ORDER BY batch")]
        self.log_subject.config(values=[ALL_SUBJECTS] + batches)
        where, args = [], []
        if self.log_status.get() != "all":
            where.append("status=?")
            args.append(self.log_status.get())
        if self.log_subject.get() != ALL_SUBJECTS:
            where.append("batch=?")
            args.append(self.log_subject.get())
        q = self.log_search.get().strip()
        if q:
            where.append("(student_id LIKE ? OR name LIKE ? OR recipients LIKE ? OR subject LIKE ?)")
            args += [f"%{q}%"] * 4
        rows = db("SELECT ts,subject,student_id,name,recipients,status,error FROM log"
                  + (" WHERE " + " AND ".join(where) if where else "") + " ORDER BY id DESC LIMIT 5000", tuple(args))
        self.log_tree.delete(*self.log_tree.get_children())
        for r in rows:
            self.log_tree.insert("", "end", values=r, tags=("failed",) if r[5] == "failed" else ())
        count = lambda st: sum(1 for r in rows if r[5] == st)
        self.log_count.set(f"{len(rows)} rows shown | sent: {count('sent')} | failed: {count('failed')} | "
                           f"partial: {count('partial')} | test: {count('test')}")

    def export_log(self):
        p = filedialog.asksaveasfilename(defaultextension=".xlsx", initialfile="send_log.xlsx", filetypes=[("Excel", "*.xlsx")])
        if not p:
            return
        wb = Workbook()
        ws = wb.active
        ws.title = "Send log"
        ws.append(["Time", "Subject", "ID", "Student", "Recipients", "Attachments", "Status", "Error / note"])
        for r in db("SELECT ts,subject,student_id,name,recipients,attachments,status,error FROM log ORDER BY id"):
            ws.append(list(r))
        try:
            wb.save(p)
            messagebox.showinfo("Exported", p)
        except OSError as e:
            messagebox.showerror("Cannot save", str(e))


if __name__ == "__main__":
    App().mainloop()
